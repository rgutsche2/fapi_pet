#%%
"""Arztbriefe auswerten: gemeinsamer Ablauf für SpA, RA und PsA.

Im bestehenden Ordner fapi_pet speichern. Prompts und Schemata liegen daneben.
In PyCharm diese Datei starten. Einstellungen direkt unten ändern.
"""
import hashlib
import json
import math
import re
import shutil
import subprocess
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ==================== EINSTELLUNGEN ====================
PROJEKTORDNER = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
INDIKATIONEN = ["SpA", "RA", "PsA"]  # Für nur PsA: ["PsA"]
PATIENTENLISTE = PROJEKTORDNER / "data" / "patients_fapi_pet.xlsx"
BRIEFORDNER = PROJEKTORDNER / "data" / "Arztbriefe_Anonymisiert"
AUSGABEORDNER = PROJEKTORDNER / "output"
PROMPTORDNER = PROJEKTORDNER / "prompts"
SCHEMAORDNER = PROJEKTORDNER / "schemata"
FOLLOW_UP_NUMMER = 1  # Erster FU bei mehreren Scans (z.B. Patienten 13, 23, 45).
NEU_AUSWERTEN = False  # Geänderte Prompts/Schemata werden automatisch neu ausgewertet.
CODEX_TIMEOUT_SEKUNDEN = 600

PHASEN = ("Baseline", "Follow-up")


def datum(wert):
    if wert is None or wert == "":
        return None
    if isinstance(wert, datetime):
        return wert.date().isoformat()
    if isinstance(wert, date):
        return wert.isoformat()
    return date.fromisoformat(str(wert).strip()[:10]).isoformat()


def konfiguration_laden(indikation):
    if indikation not in ("SpA", "RA", "PsA"):
        raise ValueError(f"Unbekannte Indikation: {indikation}")
    schema = json.loads((SCHEMAORDNER / f"{indikation}.json").read_text(encoding="utf-8"))
    if schema["indikation"] != indikation:
        raise ValueError("Indikation und Schema passen nicht zusammen.")
    spalten = []

    def hinzufuegen(name, typ, gruppe, breite):
        if not isinstance(name, str) or not name.strip() or typ not in ("Text", "Zahl", "Datum", "Formel", "Liste"):
            raise ValueError("Ungültiger Spaltenname oder Feldtyp im Schema.")
        spalten.append({"name": name, "typ": typ, "gruppe": gruppe, "breite": breite})

    for name, typ, breite in [
        ("Patient", "Zahl", 10), ("Erkrankung", "Text", 12), ("Zeitpunkt", "Text", 15),
        ("Scan-Datum", "Datum", 14), ("Klinisches Datum", "Datum", 16),
        ("Abstand zum Scan (Tage)", "Formel", 18),
    ]:
        hinzufuegen(name, typ, "Identifikation", breite)
    for feld in schema["scores"]:
        hinzufuegen(feld["name"], feld["typ"], "Scores", 18 if feld["typ"] == "Text" else 12)
    for name, typ, breite in [
        ("CRP-Wert", "Zahl", 13), ("CRP-Vergleichszeichen", "Text", 18),
        ("CRP-Einheit", "Text", 14), ("BSG (mm nach 1 h)", "Zahl", 18),
    ]:
        hinzufuegen(name, typ, "Labor", breite)
    hinzufuegen("Aktuelle Therapie", "Text", "Therapie", 55)
    for name in schema["klinische_spalten"]:
        hinzufuegen(name, "Text", "Klinische Befunde", 55)
    hinzufuegen("Verlauf", "Text", "Klinische Befunde", 65)
    hinzufuegen("Ansprechen (Ja/Nein)", "Text", "Klinische Befunde", 22)
    hinzufuegen("Hinweise", "Liste", "Klinische Befunde", 65)
    namen = [s["name"] for s in spalten]
    if len(set(namen)) != len(namen):
        raise ValueError("Doppelte Spaltennamen im Schema.")
    felder = [s["name"] for s in spalten if s["name"] not in {
        "Patient", "Erkrankung", "Zeitpunkt", "Scan-Datum", "Abstand zum Scan (Tage)", "Hinweise",
    }]
    auswahl = {**schema.get("auswahlwerte", {}),
               "CRP-Vergleichszeichen": ["=", "<", "<=", ">", ">="],
               "Ansprechen (Ja/Nein)": ["Ja", "Nein"]}
    for name in auswahl:
        if name not in felder:
            raise ValueError(f"Auswahlwerte für unbekanntes Feld: {name}")
    for paar in schema.get("gekoppelte_felder", []):
        if len(paar) != 2 or any(name not in felder for name in paar):
            raise ValueError("Gekoppelte Felder müssen zwei bekannte Feldnamen enthalten.")
    struktur = {"Zeilen": [dict(Zeitpunkt=z, **{f: None for f in felder}, Belege={}, Hinweise=[]) for z in PHASEN]}
    allgemein = (PROMPTORDNER / "allgemein.txt").read_text(encoding="utf-8")
    fachlich = (PROMPTORDNER / f"{indikation}.txt").read_text(encoding="utf-8")
    typen = {s["name"]: s["typ"] for s in spalten if s["name"] in felder}
    prompt = allgemein + "\nKRANKHEITSSPEZIFISCHE REGELN:\n" + fachlich
    prompt += "\nFELDTYPEN (fehlend bleibt null):\n" + json.dumps(typen, ensure_ascii=False)
    prompt += "\nAUSGABESTRUKTUR:\n" + json.dumps(struktur, ensure_ascii=False, indent=2)
    return {"indikation": indikation, "schema": schema, "spalten": spalten, "felder": felder,
            "zahlenfelder": [s["name"] for s in spalten if s["typ"] == "Zahl" and s["name"] in felder],
            "auswahl": auswahl, "struktur": struktur, "prompt": prompt}


def patienten_laden(indikation):
    buch = load_workbook(PATIENTENLISTE, data_only=True, read_only=True)
    try:
        zeilen = list(buch["Baseline + Follow-up"].values)
    finally:
        buch.close()
    kopf = {str(v).strip(): i for i, v in enumerate(zeilen[0]) if v is not None}
    benoetigt = {"PID", "Studiengruppe", "Scandatum", "Follow-up Daten"}
    if not benoetigt <= set(kopf):
        raise ValueError("Patientenliste: Pflichtspalten fehlen: " + ", ".join(sorted(benoetigt - set(kopf))))
    patienten = []
    for r in zeilen[1:]:
        if r[kopf["PID"]] is None or str(r[kopf["Studiengruppe"]]).strip().casefold() != indikation.casefold():
            continue
        pid = int(r[kopf["PID"]])
        baseline = datum(r[kopf["Scandatum"]])
        fu = [datum(x) for x in str(r[kopf["Follow-up Daten"]] or "").split("|") if x.strip()]
        index = FOLLOW_UP_NUMMER - 1 if len(fu) > 1 else 0
        if baseline is None or not fu or index >= len(fu):
            raise ValueError(f"Patient {pid}: Baseline oder gewählter Follow-up-Scan fehlt.")
        patienten.append({"Patient": pid, "Erkrankung": indikation,
                          "Scans": {"Baseline": baseline, "Follow-up": fu[index]}})
    if len({p["Patient"] for p in patienten}) != len(patienten):
        raise ValueError(f"Doppelte {indikation}-Patientennummer in der Patientenliste.")
    return sorted(patienten, key=lambda p: p["Patient"])


def word_lesen(pfad):
    teile = []
    for block in Document(pfad).iter_inner_content():
        if isinstance(block, Paragraph):
            teile.append(block.text)
        else:
            for zeile in block.rows:
                teile.append(" | ".join(zelle.text for zelle in zeile.cells))
    text = "\n".join(teile)
    if not text.strip():
        raise ValueError(f"Kein lesbarer Text: {pfad.name}")
    return text


def briefe_laden(pid):
    texte, gruppen = {}, {z: [] for z in PHASEN}
    dateien = sorted(p for p in BRIEFORDNER.rglob("*") if p.is_file() and p.suffix.lower() == ".docx")
    for phase, name in (("Baseline", "Baseline"), ("Follow-up", "FollowUp")):
        muster = re.compile(rf"{pid}_{name}(?:_\d+)?(?:\s*\(\d+\))?\.docx", re.I)
        gesehen = set()
        for pfad in dateien:
            if not muster.fullmatch(pfad.name):
                continue
            fingerprint = hashlib.sha256(pfad.read_bytes()).hexdigest()
            if fingerprint in gesehen:
                continue
            gesehen.add(fingerprint)
            quelle = pfad.relative_to(BRIEFORDNER).as_posix()
            texte[quelle] = word_lesen(pfad)
            gruppen[phase].append(quelle)
    if not all(gruppen.values()):
        raise ValueError("Arztbrief fehlt für " + ", ".join(k for k, v in gruppen.items() if not v))
    return texte, gruppen


def prompt_erstellen(config, patient, texte, gruppen):
    prompt = config["prompt"] + "\nPATIENT UND SCAN-DATEN:\n" + json.dumps(patient, ensure_ascii=False)
    for phase in PHASEN:
        for quelle in gruppen[phase]:
            prompt += f"\nBRIEFGRUPPE: {phase}\nDATEI: {quelle}\n{texte[quelle]}\nENDE BRIEF\n"
    # Auch Typen, erlaubte Werte und neue Spalten machen alte Ergebnisse ungültig.
    inhalt = json.dumps({"prompt": prompt, "schema": config["schema"], "spalten": config["spalten"]},
                        ensure_ascii=False, sort_keys=True)
    return prompt, hashlib.sha256(inhalt.encode("utf-8")).hexdigest()


def antwort_pruefen(daten, config, texte, gruppen):
    if not isinstance(daten, dict) or set(daten) != {"Zeilen"}:
        raise ValueError("Erwartet wird ein JSON-Objekt mit 'Zeilen'.")
    zeilen = daten["Zeilen"]
    if not isinstance(zeilen, list) or len(zeilen) != 2 or not all(isinstance(z, dict) for z in zeilen):
        raise ValueError("Erwartet werden genau zwei Zeilen als JSON-Objekte.")
    if sorted(z.get("Zeitpunkt", "") for z in zeilen if isinstance(z.get("Zeitpunkt"), str)) != sorted(PHASEN):
        raise ValueError("Baseline und Follow-up müssen jeweils einmal vorkommen.")
    for z in zeilen:
        if set(z) != set(config["felder"]) | {"Zeitpunkt", "Belege", "Hinweise"}:
            raise ValueError("Unerwartete oder fehlende JSON-Felder.")
        if not isinstance(z["Hinweise"], list) or not all(isinstance(x, str) for x in z["Hinweise"]):
            raise ValueError("Hinweise muss eine Textliste sein.")
        if not isinstance(z["Belege"], dict) or set(z["Belege"]) - set(config["felder"]):
            raise ValueError("Belege muss ein Objekt mit bekannten Feldnamen sein.")
        for f in config["felder"]:
            v = z[f]
            if f in config["zahlenfelder"]:
                if v is not None and (isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v)):
                    raise ValueError(f"{f}: Zahl oder null erwartet.")
            elif v is not None and (not isinstance(v, str) or not v.strip()):
                raise ValueError(f"{f}: nichtleerer Text oder null erwartet.")
            if f == "Klinisches Datum" and v is not None and datum(v) != v:
                raise ValueError("Klinisches Datum muss YYYY-MM-DD sein.")
            if v is not None and not (f == "Verlauf" and v == "Ausgangsbefund") and not z["Belege"].get(f):
                raise ValueError(f"Originalbeleg fehlt für {f}.")
        for f, erlaubt in config["auswahl"].items():
            if z[f] is not None and z[f] not in erlaubt:
                raise ValueError(f"Unbekannter Wert für {f}: {z[f]}")
        for a, b in config["schema"].get("gekoppelte_felder", []):
            if (z[a] is None) != (z[b] is None):
                raise ValueError(f"{a} und {b} müssen zusammen angegeben werden.")
        if z["Zeitpunkt"] == "Baseline" and z["Ansprechen (Ja/Nein)"] is not None:
            raise ValueError("Ansprechen wird nur im Follow-up beurteilt.")
        for feld, belege in z["Belege"].items():
            if not isinstance(belege, list):
                raise ValueError("Belege je Feld müssen eine Liste sein.")
            for b in belege:
                if not isinstance(b, dict) or set(b) != {"Datei", "Zitat"}:
                    raise ValueError("Ein Beleg benötigt Datei und Zitat.")
                quellen = texte if feld in ("Verlauf", "Ansprechen (Ja/Nein)") else gruppen[z["Zeitpunkt"]]
                if not isinstance(b["Datei"], str) or b["Datei"] not in quellen:
                    raise ValueError("Beleg stammt aus der falschen Briefgruppe.")
                if not isinstance(b["Zitat"], str) or not b["Zitat"].strip():
                    raise ValueError("Leeres Originalzitat.")
                if " ".join(b["Zitat"].split()) not in " ".join(texte[b["Datei"]].split()):
                    raise ValueError(f"Originalzitat nicht gefunden: {b['Zitat']}")
    return sorted(zeilen, key=lambda z: z["Zeitpunkt"] != "Baseline")


def json_speichern(ziel, daten):
    temp = ziel.with_suffix(".tmp.json")
    temp.write_text(json.dumps(daten, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(ziel)


def excel_speichern(config, zeilen, ordner):
    buch = Workbook()
    blatt = buch.active
    blatt.title = config["indikation"]
    spalten = config["spalten"]
    index = {s["name"]: i for i, s in enumerate(spalten, 1)}
    farben = {"Identifikation": "D9EAF7", "Scores": "E2EFDA", "Labor": "FFF2CC",
              "Therapie": "FCE4D6", "Klinische Befunde": "CBEAF5"}
    rand = Border(bottom=Side(style="thin", color="D9E2F3"))
    start = 1
    for j, s in enumerate(spalten, 1):
        gruppe = s["gruppe"]
        blatt.column_dimensions[get_column_letter(j)].width = s["breite"]
        c = blatt.cell(2, j, s["name"])
        c.font = Font(name="Calibri", size=11, bold=True)
        c.fill = PatternFill("solid", fgColor=farben[gruppe])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if j == len(spalten) or spalten[j]["gruppe"] != gruppe:
            if start < j:
                blatt.merge_cells(start_row=1, start_column=start, end_row=1, end_column=j)
            c = blatt.cell(1, start, gruppe)
            c.font = Font(name="Calibri", size=12, bold=True)
            c.fill = PatternFill("solid", fgColor=farben[gruppe])
            c.alignment = Alignment(vertical="center")
            start = j + 1
    for i, z in enumerate(zeilen, 3):
        max_zeilen = 2
        for j, s in enumerate(spalten, 1):
            name = s["name"]
            wert = z.get(name)
            if name == "Hinweise":
                wert = "\n".join(wert or []) or None
            if s["typ"] == "Datum" and wert:
                wert = date.fromisoformat(wert)
            c = blatt.cell(i, j, wert)
            c.font = Font(name="Calibri", size=11)
            c.border = rand
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.fill = PatternFill("solid", fgColor="F2F6FA" if z["Zeitpunkt"] == "Baseline" else "FFFFFF")
            if s["typ"] == "Datum":
                c.number_format = "yyyy-mm-dd"
            elif s["typ"] in ("Zahl", "Formel"):
                c.number_format = "0" if name in ("Patient", "Abstand zum Scan (Tage)") else "0.###"
            if isinstance(wert, str):
                c.data_type = "s"  # Freitext bleibt auch bei führendem '=' normaler Text.
                breite = max(8, s["breite"] - 3)
                max_zeilen = max(max_zeilen, sum(max(1, math.ceil(len(t) / breite)) for t in wert.split("\n")))
        klinisch = f"{get_column_letter(index['Klinisches Datum'])}{i}"
        scan = f"{get_column_letter(index['Scan-Datum'])}{i}"
        blatt.cell(i, index["Abstand zum Scan (Tage)"], f'=IF({klinisch}="","",{klinisch}-{scan})')
        blatt.row_dimensions[i].height = min(409, max(45, max_zeilen * 15))
    ende = len(zeilen) + 2
    for name, werte in config["auswahl"].items():
        validierung = DataValidation(type="list", formula1='"' + ",".join(werte) + '"', allow_blank=True)
        blatt.add_data_validation(validierung)
        validierung.add(f"{get_column_letter(index[name])}3:{get_column_letter(index[name])}{max(3, ende)}")
    blatt.row_dimensions[1].height = 24
    blatt.row_dimensions[2].height = 46
    blatt.freeze_panes = "D3"
    blatt.auto_filter.ref = f"A2:{get_column_letter(len(spalten))}{ende}"
    ziel = ordner / f"{config['indikation']}_Auswertung.xlsx"
    temp = ziel.with_suffix(".tmp.xlsx")
    try:
        buch.save(temp)
        temp.replace(ziel)
    finally:
        buch.close()
    return ziel


def auswerten(config, patienten, codex):
    indikation = config["indikation"]
    ordner = AUSGABEORDNER / indikation
    json_ordner = ordner / "json"
    json_ordner.mkdir(parents=True, exist_ok=True)
    alle_zeilen = [{"Patient": p["Patient"], "Erkrankung": indikation, "Zeitpunkt": z,
                    "Scan-Datum": p["Scans"][z], "Verlauf": "Auswertung ausstehend"}
                   for p in patienten for z in PHASEN]
    print(f"\n{indikation}: {len(patienten)} Patienten, {len(alle_zeilen)} Zeilen.", flush=True)
    erfolgreich, stoppen = 0, False
    for nummer, p in enumerate(patienten):
        print(f"[{nummer + 1}/{len(patienten)}] Patient {p['Patient']}", flush=True)
        try:
            texte, gruppen = briefe_laden(p["Patient"])
            prompt, fingerprint = prompt_erstellen(config, p, texte, gruppen)
            ziel = json_ordner / f"{p['Patient']}_{indikation}.json"
            daten = None
            if ziel.exists() and not NEU_AUSWERTEN:
                try:
                    gespeichert = json.loads(ziel.read_text(encoding="utf-8"))
                    if isinstance(gespeichert, dict) and gespeichert.get("fingerprint") == fingerprint:
                        daten = gespeichert["Ergebnis"]
                        print("Unverändertes vorhandenes Ergebnis wird verwendet.")
                except (ValueError, KeyError):
                    print("Vorhandenes JSON nicht lesbar; Patient wird erneut ausgewertet.")
            if daten is None:
                print("Warte auf Codex ...", flush=True)
                antwort = subprocess.run(
                    [codex, "exec", "--skip-git-repo-check", "--sandbox", "read-only", "-"],
                    input=prompt, capture_output=True, text=True, encoding="utf-8",
                    timeout=CODEX_TIMEOUT_SEKUNDEN, cwd=PROJEKTORDNER,
                )
                if antwort.returncode != 0:
                    raise RuntimeError(antwort.stderr[-2000:] or "Codex ohne gültige Antwort beendet.")
                (json_ordner / f"{p['Patient']}_{indikation}_roh.txt").write_text(antwort.stdout, encoding="utf-8")
                daten = json.loads(antwort.stdout)
            zeilen = antwort_pruefen(daten, config, texte, gruppen)
            json_speichern(ziel, {"Patient": p, "fingerprint": fingerprint, "Ergebnis": daten})
            for offset, z in enumerate(zeilen):
                alle_zeilen[2 * nummer + offset].update(z)
                if z["Hinweise"]:
                    print(z["Zeitpunkt"] + " – Hinweise: " + " | ".join(z["Hinweise"]))
            erfolgreich += 1
            print("JSON gespeichert.", flush=True)
        except KeyboardInterrupt:
            print("Unterbrochen; gespeicherte Ergebnisse bleiben erhalten.")
            stoppen = True
        except Exception as fehler:
            meldung = f"{type(fehler).__name__}: {fehler}"
            print("Fehler: " + meldung, flush=True)
            for offset in (0, 1):
                alle_zeilen[2 * nummer + offset]["Verlauf"] = "Keine Auswertung: " + meldung
            if isinstance(fehler, (RuntimeError, subprocess.TimeoutExpired)):
                print("Codex-Lauf gestoppt. Erneut starten, um gespeicherte Ergebnisse weiterzuverwenden.")
                stoppen = True
        finally:
            excel_speichern(config, alle_zeilen, ordner)
        if stoppen:
            break
    print(f"Fertig ausgewertet: {erfolgreich}/{len(patienten)} Patienten.")
    print(f"Excel: {ordner / (indikation + '_Auswertung.xlsx')}")
    print(f"JSON und Originalbelege: {json_ordner}")
    return stoppen


def main():
    if not PATIENTENLISTE.is_file() or not BRIEFORDNER.is_dir():
        raise SystemExit(f"Patientenliste oder Briefordner fehlt. Prüfe {PATIENTENLISTE} und {BRIEFORDNER}.")
    if not INDIKATIONEN or len(set(INDIKATIONEN)) != len(INDIKATIONEN):
        raise SystemExit("INDIKATIONEN darf nicht leer sein oder doppelte Einträge enthalten.")
    if isinstance(FOLLOW_UP_NUMMER, bool) or not isinstance(FOLLOW_UP_NUMMER, int) or FOLLOW_UP_NUMMER < 1:
        raise SystemExit("FOLLOW_UP_NUMMER muss eine positive ganze Zahl sein.")
    # Alle Konfigurationen vor dem ersten Modellaufruf prüfen.
    konfigurationen = [konfiguration_laden(i) for i in INDIKATIONEN]
    listen = {c["indikation"]: patienten_laden(c["indikation"]) for c in konfigurationen}
    codex = shutil.which("codex")
    if codex is None:
        raise SystemExit("Codex nicht gefunden. Nach der Installation PyCharm neu starten.")
    for config in konfigurationen:
        patienten = listen[config["indikation"]]
        if not patienten:
            print(f"Keine {config['indikation']}-Patienten in der Liste; übersprungen.")
            continue
        if auswerten(config, patienten, codex):
            break


if __name__ == "__main__":
    main()
