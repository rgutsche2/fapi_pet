"""
Erzeugt formatierte Excel-Ausgaben mit openpyxl.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def safe_sheet_name(name: str) -> str:
    """Kürzt Blattnamen auf das von Excel erlaubte Maximum."""
    return name[:31]


def write_excel_file(
    output_file: Path,
    sheets: dict[str, pd.DataFrame],
) -> None:
    """Schreibt mehrere DataFrames in eine formatierte Excel-Datei."""
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl",
        date_format="DD.MM.YYYY",
        datetime_format="DD.MM.YYYY",
    ) as writer:
        for sheet_name, data in sheets.items():
            export_data = data.copy()

            # Interne Hilfsspalten nicht exportieren
            internal_columns = [
                column for column in export_data.columns if column.startswith("_")
            ]
            export_data = export_data.drop(columns=internal_columns, errors="ignore")

            export_data.to_excel(
                writer,
                sheet_name=safe_sheet_name(sheet_name),
                index=False,
            )

    format_workbook(output_file)


def format_workbook(output_file: Path) -> None:
    """Wendet Filter, Tabellenformat, Freeze-Panes und Spaltenbreiten an."""
    workbook = load_workbook(output_file)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        worksheet.row_dimensions[1].height = 24

        if worksheet.max_row >= 2 and worksheet.max_column >= 1:
            table_ref = (
                f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
            )
            table_name = (
                "Tab_"
                + "".join(character for character in worksheet.title if character.isalnum())
            )[:250]

            table = Table(displayName=table_name, ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            worksheet.add_table(table)

        for column_index, column_cells in enumerate(
            worksheet.iter_cols(min_row=1, max_row=worksheet.max_row),
            start=1,
        ):
            max_length = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

                if "datum" in str(worksheet.cell(1, column_index).value).lower():
                    if cell.row > 1 and cell.value is not None:
                        cell.number_format = "DD.MM.YYYY"

                cell.alignment = Alignment(vertical="top", wrap_text=True)

            width = min(max(max_length + 2, 12), 55)
            worksheet.column_dimensions[get_column_letter(column_index)].width = width

    workbook.save(output_file)
